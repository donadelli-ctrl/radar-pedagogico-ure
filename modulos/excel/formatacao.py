"""
==========================================================
RADAR PEDAGÓGICO URE
MÓDULO: formatacao.py
Versão: 3.0
==========================================================

Responsabilidade:
Centralizar a formatação visual das planilhas Excel.

Esta função NÃO altera:
- dados;
- indicadores;
- critérios;
- classificação do Farol;
- diagnóstico.

Apenas aplica a identidade visual do relatório.
"""

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

from openpyxl.utils import get_column_letter


# ==========================================================
# PALETA VISUAL
# ==========================================================

# Cabeçalho único da tabela
COR_CABECALHO = "C9E2EE"

# Texto
COR_TEXTO = "111111"
COR_TEXTO_FORTE = "000000"


# ==========================================================
# CORES DOS INDICADORES
# ==========================================================

# Vermelho — desempenho abaixo do esperado
COR_VERMELHO = "F4A3A3"

# Amarelo — atenção
COR_AMARELO = "FFD966"

# Verde — resultado favorável
COR_VERDE = "A9D18E"

# Azul — destaque
COR_AZUL = "8FBCE6"

# Cinza — sem dados
COR_CINZA = "D9D9D9"


# ==========================================================
# CORES DA SITUAÇÃO DO FAROL
# ==========================================================

COR_PRIORITARIA = PatternFill(
    fill_type="solid",
    fgColor="F4A3A3",
)

COR_ATENCAO = PatternFill(
    fill_type="solid",
    fgColor="FFD966",
)

COR_FAVORAVEL = PatternFill(
    fill_type="solid",
    fgColor="A9D18E",
)

COR_DESTAQUE = PatternFill(
    fill_type="solid",
    fgColor="8FBCE6",
)

COR_SEM_DADOS = PatternFill(
    fill_type="solid",
    fgColor="D9D9D9",
)


# ==========================================================
# PREENCHIMENTO DO CABEÇALHO
# ==========================================================

CABECALHO_FILL = PatternFill(
    fill_type="solid",
    fgColor=COR_CABECALHO,
)


# ==========================================================
# FONTE DO CABEÇALHO
# ==========================================================

CABECALHO_FONT = Font(
    name="Aptos",
    size=10,
    bold=True,
    color=COR_TEXTO_FORTE,
)


# ==========================================================
# ALINHAMENTOS
# ==========================================================

CENTRO = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)

ESQUERDA = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=False,
)

ESQUERDA_COMPACTA = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True,
)

# ==========================================================
# BORDA
# ==========================================================

BORDA = Border(
    left=Side(
        style="thin",
        color="D0D0D0",
    ),
    right=Side(
        style="thin",
        color="D0D0D0",
    ),
    top=Side(
        style="thin",
        color="D0D0D0",
    ),
    bottom=Side(
        style="thin",
        color="D0D0D0",
    ),
)


# ==========================================================
# FORMATAR PLANILHA
# ==========================================================

def formatar_planilha(ws):
    """
    Aplica a identidade visual padrão do Radar.

    A primeira linha possui uma única cor.
    As cores dos indicadores aparecem somente
    no corpo da tabela.
    """

    # ======================================================
    # CABEÇALHO
    # ======================================================

    for cell in ws[1]:

        # ----------------------------------------------
        # UMA ÚNICA COR NO CABEÇALHO
        # ----------------------------------------------

        cell.fill = CABECALHO_FILL

        # ----------------------------------------------
        # FONTE PRETA E FORTE
        # ----------------------------------------------

        cell.font = CABECALHO_FONT

        # ----------------------------------------------
        # ALINHAMENTO
        # ----------------------------------------------

        cell.alignment = CENTRO

        # ----------------------------------------------
        # BORDA
        # ----------------------------------------------

        cell.border = BORDA

    # ======================================================
    # ALTURA DO CABEÇALHO
    # ======================================================

    ws.row_dimensions[1].height = 25

    # ======================================================
    # PERCENTUAIS
    # ======================================================

    PREFIXOS_PERCENTUAIS = (
        "PART_",
        "MEDIA_",
        "LP_",
        "MAT_",
    )

    # ======================================================
    # FORMATAÇÃO DAS COLUNAS
    # ======================================================

    for coluna in ws.columns:

        letra = get_column_letter(
            coluna[0].column
        )

        cabecalho = str(
            coluna[0].value
        ).strip().upper()

        maior = len(cabecalho)

        # --------------------------------------------------
        # CORPO
        # --------------------------------------------------

        for cell in coluna:

            if cell.row > 1:

                # ------------------------------------------
                # ESCOLA
                # ------------------------------------------

                if cabecalho == "ESCOLA":

                    cell.alignment = ESQUERDA

                    cell.font = Font(
                        name="Aptos",
                        size=9,
                        bold=True,
                        color=COR_TEXTO_FORTE,
                    )

                # ------------------------------------------
                # TEXTOS LONGOS
                # ------------------------------------------

                elif cabecalho in (
                    "DIAGNOSTICO",
                    "ENCAMINHAMENTO",
                    "COMPONENTE_FAROL",
                ):

                    cell.alignment = ESQUERDA_COMPACTA

                    cell.font = Font(
                        name="Aptos",
                        size=9,
                        color=COR_TEXTO,
                    )

                # ------------------------------------------
                # DEMAIS CAMPOS
                # ------------------------------------------

                else:

                    cell.alignment = CENTRO

                    cell.font = Font(
                        name="Aptos",
                        size=9,
                        color=COR_TEXTO,
                    )

                # ------------------------------------------
                # FORMATO DE PERCENTUAL
                # ------------------------------------------

                if (
                    cabecalho.startswith(
                        PREFIXOS_PERCENTUAIS
                    )
                    and isinstance(
                        cell.value,
                        (int, float),
                    )
                ):

                    cell.number_format = "0.0%"

                # ------------------------------------------
                # BORDA
                # ------------------------------------------

                cell.border = BORDA

            # --------------------------------------------------
            # TAMANHO DO CONTEÚDO
            # --------------------------------------------------

            valor = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            if len(valor) > maior:

                maior = len(valor)

        # ==================================================
        # LARGURA DAS COLUNAS
        # ==================================================

        if cabecalho == "ESCOLA":

            largura = 34

        elif cabecalho == "CIE":

            largura = 9

        elif cabecalho.startswith("PART_"):

            largura = 11

        elif cabecalho.startswith("MEDIA_"):

            largura = 10

        elif cabecalho.startswith("LP_"):

            largura = 10

        elif cabecalho.startswith("MAT_"):

            largura = 10

        elif cabecalho == "FAROL":

            largura = 7

        elif cabecalho == "FAROL_URE":

            largura = 16

        elif cabecalho == "COMPONENTE_FAROL":

            largura = 23

        elif cabecalho == "DESTAQUE_EVOLUCAO":

            largura = 17

        elif cabecalho == "DIAGNOSTICO":

            largura = 48

        elif cabecalho == "ENCAMINHAMENTO":

            largura = 52

        elif cabecalho == "SITUACAO":

            largura = 16

        else:

            largura = min(
                maior + 2,
                28,
            )

        ws.column_dimensions[
            letra
        ].width = largura

    # ======================================================
    # ALTURA DAS LINHAS
    # ======================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        ws.row_dimensions[
            linha
        ].height = 25

    # ======================================================
    # SITUAÇÃO DO FAROL
    # ======================================================

    for coluna in ws.columns:

        cabecalho = str(
            coluna[0].value
        ).strip().upper()

        if cabecalho not in (
            "SITUACAO",
            "FAROL_URE",
        ):

            continue

        for cell in coluna[1:]:

            valor = str(
                cell.value
            ).upper()

            if "PRIORIT" in valor:

                cell.fill = COR_PRIORITARIA

                cell.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="8B0000",
                )

            elif "ATEN" in valor:

                cell.fill = COR_ATENCAO

                cell.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="7F6000",
                )

            elif "FAVOR" in valor:

                cell.fill = COR_FAVORAVEL

                cell.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="27632A",
                )

            elif "DESTAQUE" in valor:

                cell.fill = COR_DESTAQUE

                cell.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="1F4E79",
                )

            elif "SEM" in valor:

                cell.fill = COR_SEM_DADOS

    # ======================================================
    # CORES DOS INDICADORES
    # ======================================================

    for coluna in ws.columns:

        cabecalho = str(
            coluna[0].value
        ).strip().upper()

        if not (
            cabecalho.startswith("LP")
            or cabecalho.startswith("MAT")
            or cabecalho.startswith("PART")
        ):

            continue

        for cell in coluna[1:]:

            if not isinstance(
                cell.value,
                (int, float),
            ):

                continue

            valor = cell.value

            # --------------------------------------------------
            # Decimal → percentual
            # --------------------------------------------------

            if valor <= 1:

                valor *= 100

            # ==================================================
            # PARTICIPAÇÃO
            # ==================================================

            if cabecalho.startswith(
                "PART"
            ):

                if valor < 90:

                    cell.fill = COR_PRIORITARIA

                elif valor < 95:

                    cell.fill = COR_ATENCAO

                else:

                    cell.fill = COR_FAVORAVEL

            # ==================================================
            # DESEMPENHO
            # ==================================================

            else:

                if valor < 50:

                    cell.fill = COR_PRIORITARIA

                elif valor < 70:

                    cell.fill = COR_ATENCAO

                elif valor < 90:

                    cell.fill = COR_FAVORAVEL

                else:

                    cell.fill = COR_DESTAQUE

    # ======================================================
    # CONGELAMENTO
    # ======================================================
    # ======================================================
    # FILTRO
    # ======================================================

    if ws.max_row >= 2:

        ws.auto_filter.ref = ws.dimensions