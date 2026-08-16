"""
==========================================================
RADAR PEDAGÓGICO URE
MÓDULO: prioritarias.py
Versão: 3.0
==========================================================

Responsabilidade:
Gerar a aba ESCOLAS PRIORITÁRIAS.

A classificação é determinada exclusivamente pelo
FAROL_URE calculado pelo Radar Pedagógico URE.

O Farol SARESP da PP2 não participa da classificação
e não aparece no relatório final.
"""

import pandas as pd

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)

from modulos.excel.formatacao import (
    formatar_planilha,
)


# ==========================================================
# CORES
# ==========================================================

FILL_PRIORITARIA = PatternFill(
    fill_type="solid",
    fgColor="F4A3A3",
)

FILL_DESTAQUE = PatternFill(
    fill_type="solid",
    fgColor="8FBCE6",
)


# ==========================================================
# FONTES
# ==========================================================

FONTE_NORMAL = Font(
    name="Aptos",
    size=9,
    color="111111",
)

FONTE_ESCOLA = Font(
    name="Aptos",
    size=9,
    bold=True,
    color="000000",
)

FONTE_CABECALHO = Font(
    name="Aptos",
    size=10,
    bold=True,
    color="000000",
)

FONTE_PRIORITARIA = Font(
    name="Aptos",
    size=9,
    bold=True,
    color="9C0006",
)


# ==========================================================
# BORDA
# ==========================================================

BORDA = Border(
    left=Side(
        style="thin",
        color="D9D9D9",
    ),
    right=Side(
        style="thin",
        color="D9D9D9",
    ),
    top=Side(
        style="thin",
        color="D9D9D9",
    ),
    bottom=Side(
        style="thin",
        color="D9D9D9",
    ),
)


# ==========================================================
# ORDEM DAS COLUNAS
# ==========================================================

COLUNAS_PRIORITARIAS = [
    "ESCOLA",

    # ------------------------------------------------------
    # DESEMPENHO ADE
    # ------------------------------------------------------

    "LP_ABAIXO",
    "MAT_ABAIXO",

    # ------------------------------------------------------
    # PROVA PAULISTA
    # ------------------------------------------------------

    "LP_PP1",
    "MAT_PP1",

    "LP_PP2",
    "MAT_PP2",

    # ------------------------------------------------------
    # PARTICIPAÇÃO
    # ------------------------------------------------------

    "PART_ADE",
    "PART_PP1",
    "PART_PP2",

    # ------------------------------------------------------
    # EVOLUÇÃO
    # ------------------------------------------------------

    "EVOLUCAO_LP",
    "EVOLUCAO_MAT",

    # ------------------------------------------------------
    # FAROL
    # ------------------------------------------------------

    "FAROL",
    "FAROL_URE",
    "COMPONENTE_FAROL",
    "DESTAQUE_EVOLUCAO",

    # ------------------------------------------------------
    # ANÁLISE PEDAGÓGICA
    # ------------------------------------------------------

    "DIAGNOSTICO",
    "ENCAMINHAMENTO",
]


# ==========================================================
# LOCALIZA COLUNA
# ==========================================================

def obter_indice_coluna(
    ws,
    nome_coluna,
):
    """
    Retorna o índice da coluna pelo nome.
    """

    for coluna in ws.iter_cols(
        min_row=1,
        max_row=1,
    ):

        if coluna[0].value == nome_coluna:

            return coluna[0].column

    return None


# ==========================================================
# LARGURA
# ==========================================================

def definir_largura_coluna(
    ws,
    nome_coluna,
    largura,
):
    """
    Define a largura da coluna.
    """

    indice = obter_indice_coluna(
        ws,
        nome_coluna,
    )

    if indice is None:

        return

    letra = ws.cell(
        row=1,
        column=indice,
    ).column_letter

    ws.column_dimensions[
        letra
    ].width = largura


# ==========================================================
# TEXTO LONGO
# ==========================================================

def formatar_texto_longo(
    ws,
    nome_coluna,
):
    """
    Configura as colunas de texto pedagógico.
    """

    indice = obter_indice_coluna(
        ws,
        nome_coluna,
    )

    if indice is None:

        return

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        celula = ws.cell(
            row=linha,
            column=indice,
        )

        celula.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )

        celula.font = FONTE_NORMAL


# ==========================================================
# ESCOLAS PRIORITÁRIAS
# ==========================================================

def criar_prioritarias(
    writer,
    df: pd.DataFrame,
):
    """
    Cria a aba ESCOLAS PRIORITÁRIAS.

    Somente escolas classificadas pelo FAROL_URE
    como PRIORITÁRIA são apresentadas.
    """

    if df is None or df.empty:

        return

    # ======================================================
    # FAROL OFICIAL
    # ======================================================

    if "FAROL_URE" not in df.columns:

        raise ValueError(
            "A coluna FAROL_URE não foi encontrada. "
            "Não é possível gerar a aba "
            "ESCOLAS PRIORITÁRIAS."
        )

    # ======================================================
    # SELECIONA PRIORITÁRIAS
    # ======================================================

    resultado = df[
        df["FAROL_URE"]
        .astype(str)
        .str.upper()
        .str.contains(
            "PRIORIT",
            na=False,
        )
    ].copy()

    if resultado.empty:

        return

    # ======================================================
    # SELEÇÃO DAS COLUNAS
    # ======================================================

    colunas_existentes = [
        coluna
        for coluna in COLUNAS_PRIORITARIAS
        if coluna in resultado.columns
    ]

    resultado = resultado[
        colunas_existentes
    ].copy()

    # ======================================================
    # FAROL VISUAL
    # ======================================================

    if "FAROL_URE" in resultado.columns:

        indice_farol_ure = (
            resultado.columns.get_loc(
                "FAROL_URE"
            )
        )

        resultado.insert(
            indice_farol_ure,
            "FAROL",
            "●",
        )

    # ======================================================
    # ORDENAÇÃO
    # ======================================================

    colunas_ordenacao = []

    if "COMPONENTE_FAROL" in resultado.columns:

        colunas_ordenacao.append(
            "COMPONENTE_FAROL"
        )

    if "ESCOLA" in resultado.columns:

        colunas_ordenacao.append(
            "ESCOLA"
        )

    if colunas_ordenacao:

        resultado.sort_values(
            by=colunas_ordenacao,
            ascending=True,
            inplace=True,
            kind="stable",
        )

    # ======================================================
    # EXPORTAÇÃO
    # ======================================================

    resultado.to_excel(
        writer,
        sheet_name="ESCOLAS PRIORITÁRIAS",
        index=False,
    )

    ws = writer.sheets[
        "ESCOLAS PRIORITÁRIAS"
    ]

    # ======================================================
    # FORMATAÇÃO BASE
    # ======================================================

    formatar_planilha(
        ws
    )

    # ======================================================
    # CONGELA CABEÇALHO
    # ======================================================

    ws.freeze_panes = "A2"

    # ======================================================
    # FILTRO
    # ======================================================

    if ws.max_row >= 2:

        ws.auto_filter.ref = ws.dimensions

    # ======================================================
    # CABEÇALHO
    # ======================================================

    for celula in ws[1]:

        celula.font = FONTE_CABECALHO

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        celula.border = BORDA

    # ======================================================
    # CORPO
    # ======================================================

    indice_escola = obter_indice_coluna(
        ws,
        "ESCOLA",
    )

    indice_farol = obter_indice_coluna(
        ws,
        "FAROL",
    )

    indice_farol_ure = obter_indice_coluna(
        ws,
        "FAROL_URE",
    )

    indice_diagnostico = obter_indice_coluna(
        ws,
        "DIAGNOSTICO",
    )

    indice_encaminhamento = obter_indice_coluna(
        ws,
        "ENCAMINHAMENTO",
    )

    indice_componente = obter_indice_coluna(
        ws,
        "COMPONENTE_FAROL",
    )

    indice_destaque = obter_indice_coluna(
        ws,
        "DESTAQUE_EVOLUCAO",
    )

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        # --------------------------------------------------
        # FORMATAÇÃO GERAL
        # --------------------------------------------------

        for coluna in range(
            1,
            ws.max_column + 1,
        ):

            celula = ws.cell(
                row=linha,
                column=coluna,
            )

            celula.font = FONTE_NORMAL

            celula.border = BORDA

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=False,
            )

        # --------------------------------------------------
        # ESCOLA
        # --------------------------------------------------

        if indice_escola is not None:

            celula = ws.cell(
                row=linha,
                column=indice_escola,
            )

            celula.font = FONTE_ESCOLA

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=False,
            )

        # --------------------------------------------------
        # DIAGNÓSTICO
        # --------------------------------------------------

        if indice_diagnostico is not None:

            celula = ws.cell(
                row=linha,
                column=indice_diagnostico,
            )

            celula.font = FONTE_NORMAL

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # ENCAMINHAMENTO
        # --------------------------------------------------

        if indice_encaminhamento is not None:

            celula = ws.cell(
                row=linha,
                column=indice_encaminhamento,
            )

            celula.font = FONTE_NORMAL

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # COMPONENTE DO FAROL
        # --------------------------------------------------

        if indice_componente is not None:

            celula = ws.cell(
                row=linha,
                column=indice_componente,
            )

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # FAROL
        # --------------------------------------------------

        if indice_farol is not None:

            celula = ws.cell(
                row=linha,
                column=indice_farol,
            )

            celula.fill = FILL_PRIORITARIA

            celula.font = Font(
                name="Arial",
                size=15,
                bold=True,
                color="9C0006",
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # --------------------------------------------------
        # FAROL URE
        # --------------------------------------------------

        if indice_farol_ure is not None:

            celula = ws.cell(
                row=linha,
                column=indice_farol_ure,
            )

            celula.fill = FILL_PRIORITARIA

            celula.font = FONTE_PRIORITARIA

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # DESTAQUE DE EVOLUÇÃO
        # --------------------------------------------------

        if indice_destaque is not None:

            valor = ws.cell(
                row=linha,
                column=indice_destaque,
            ).value

            if valor is True:

                ws.cell(
                    row=linha,
                    column=indice_destaque,
                ).fill = FILL_DESTAQUE

    # ======================================================
    # PERCENTUAIS
    # ======================================================

    colunas_percentuais = [
        "LP_ABAIXO",
        "MAT_ABAIXO",
        "LP_PP1",
        "MAT_PP1",
        "LP_PP2",
        "MAT_PP2",
        "PART_ADE",
        "PART_PP1",
        "PART_PP2",
        "EVOLUCAO_LP",
        "EVOLUCAO_MAT",
    ]

    for nome_coluna in colunas_percentuais:

        indice = obter_indice_coluna(
            ws,
            nome_coluna,
        )

        if indice is None:

            continue

        for linha in range(
            2,
            ws.max_row + 1,
        ):

            ws.cell(
                row=linha,
                column=indice,
            ).number_format = "0.0%"

    # ======================================================
    # LARGURAS
    # ======================================================

    definir_largura_coluna(
        ws,
        "ESCOLA",
        34,
    )

    definir_largura_coluna(
        ws,
        "FAROL",
        7,
    )

    definir_largura_coluna(
        ws,
        "FAROL_URE",
        16,
    )

    definir_largura_coluna(
        ws,
        "COMPONENTE_FAROL",
        22,
    )

    definir_largura_coluna(
        ws,
        "DESTAQUE_EVOLUCAO",
        17,
    )

    definir_largura_coluna(
        ws,
        "DIAGNOSTICO",
        48,
    )

    definir_largura_coluna(
        ws,
        "ENCAMINHAMENTO",
        52,
    )

    # ======================================================
    # TEXTOS LONGOS
    # ======================================================

    formatar_texto_longo(
        ws,
        "DIAGNOSTICO",
    )

    formatar_texto_longo(
        ws,
        "ENCAMINHAMENTO",
    )

    # ======================================================
    # ALTURA DAS LINHAS
    # ======================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        texto = ""

        if "DIAGNOSTICO" in resultado.columns:

            valor = resultado.iloc[
                linha - 2
            ]["DIAGNOSTICO"]

            if pd.notna(valor):

                texto += str(valor)

        if "ENCAMINHAMENTO" in resultado.columns:

            valor = resultado.iloc[
                linha - 2
            ]["ENCAMINHAMENTO"]

            if pd.notna(valor):

                texto += str(valor)

        tamanho = len(texto)

        if tamanho > 250:

            altura = 42

        elif tamanho > 160:

            altura = 36

        elif tamanho > 90:

            altura = 32

        else:

            altura = 25

        ws.row_dimensions[
            linha
        ].height = altura

    # ======================================================
    # CABEÇALHO
    # ======================================================

    ws.row_dimensions[
        1
    ].height = 25

    # ======================================================
    # ZOOM
    # ======================================================

    ws.sheet_view.zoomScale = 90

    # ======================================================
    # GRIDLINES
    # ======================================================

    ws.sheet_view.showGridLines = False