"""
==========================================================
RADAR PEDAGÓGICO URE
MÓDULO: radar.py
Versão: 3.4
==========================================================

Responsabilidade:
Gerar a aba RADAR PEDAGÓGICO.

Esta versão trabalha a apresentação visual do relatório.

Não altera:
- cálculos;
- indicadores;
- critérios do Farol;
- diagnóstico;
- encaminhamentos.

O Farol utilizado é exclusivamente o FAROL_URE.
"""

import pandas as pd

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)

from modulos.excel.formatacao import (
    formatar_planilha,
)


# ==========================================================
# CORES DOS GRUPOS DE AVALIAÇÃO
# ==========================================================

COR_ADE = "DDEBF7"

COR_PP1 = "B4C7E7"

COR_PP2 = "8EAADB"

COR_PARTICIPACAO = "D6E4F0"

COR_EVOLUCAO = "B4C7E7"

COR_GERAL = "D9EAF4"

COR_FUNDO_DIAGNOSTICO = "FFFFFF"
COR_FUNDO_ENCAMINHAMENTO = "F3F8FC"


# ==========================================================
# CORES PEDAGÓGICAS
# ==========================================================

FILL_PRIORITARIA = PatternFill(
    fill_type="solid",
    fgColor="F4A3A3",
)

FILL_ATENCAO = PatternFill(
    fill_type="solid",
    fgColor="FFD966",
)

FILL_FAVORAVEL = PatternFill(
    fill_type="solid",
    fgColor="A9D18E",
)

FILL_SEM_HISTORICO = PatternFill(
    fill_type="solid",
    fgColor="D9D9D9",
)

FILL_DESTAQUE = PatternFill(
    fill_type="solid",
    fgColor="8FBCE6",
)


# ==========================================================
# CORES DA EVOLUÇÃO
# ==========================================================

COR_EVOLUCAO_POSITIVA = "008000"

COR_EVOLUCAO_NEGATIVA = "C00000"

COR_EVOLUCAO_NEUTRA = "000000"


# ==========================================================
# CORES DO FAROL
# ==========================================================

COR_PRIORITARIA = "8B0000"

COR_ATENCAO = "6B5200"

COR_FAVORAVEL = "27632A"

COR_NEUTRO = "444444"


# ==========================================================
# FONTES
# ==========================================================

FONTE_NORMAL = Font(
    name="Aptos",
    size=9,
    color="000000",
)

FONTE_ESCOLA = Font(
    name="Aptos",
    size=9,
    bold=True,
    color="000000",
)

FONTE_CABECALHO = Font(
    name="Aptos",
    size=10,
    bold=True,
    color="000000",
)

FONTE_TEXTO_LONGO = Font(
    name="Aptos",
    size=9,
    color="000000",
)


# ==========================================================
# BORDAS
# ==========================================================

BORDA = Border(
    left=Side(
        style="thin",
        color="D9D9D9",
    ),
    right=Side(
        style="thin",
        color="D9D9D9",
    ),
    top=Side(
        style="thin",
        color="D9D9D9",
    ),
    bottom=Side(
        style="thin",
        color="D9D9D9",
    ),
)

BORDA_SEPARACAO_ESCOLA = Border(
    left=Side(
        style="thin",
        color="D9D9D9",
    ),
    right=Side(
        style="thin",
        color="D9D9D9",
    ),
    top=Side(
        style="thin",
        color="D9D9D9",
    ),
    bottom=Side(
        style="medium",
        color="000000",
    ),
)


# ==========================================================
# ORDEM DAS COLUNAS
# ==========================================================

COLUNAS_RADAR = [
    "ESCOLA",

    "LP_ABAIXO",
    "MAT_ABAIXO",

    "LP_PP1",
    "MAT_PP1",

    "LP_PP2",
    "MAT_PP2",

    "PART_ADE",
    "PART_PP1",
    "PART_PP2",

    "EVOLUCAO_LP",
    "EVOLUCAO_MAT",

    "FAROL",
    "FAROL_URE",

    "COMPONENTE_FAROL",

    "DESTAQUE_EVOLUCAO",

    "DIAGNOSTICO",
    "ENCAMINHAMENTO",
]


# ==========================================================
# LOCALIZA COLUNA
# ==========================================================

def obter_indice_coluna(
    ws,
    nome_coluna,
):

    for coluna in ws.iter_cols(
        min_row=1,
        max_row=1,
    ):

        if coluna[0].value == nome_coluna:

            return coluna[0].column

    return None


# ==========================================================
# LARGURA
# ==========================================================

def definir_largura_coluna(
    ws,
    nome_coluna,
    largura,
):

    indice = obter_indice_coluna(
        ws,
        nome_coluna,
    )

    if indice is None:

        return

    letra = ws.cell(
        row=1,
        column=indice,
    ).column_letter

    ws.column_dimensions[
        letra
    ].width = largura


# ==========================================================
# FAROL VISUAL
# ==========================================================

def gerar_farol(situacao):

    return "●"


# ==========================================================
# COR DO FAROL
# ==========================================================

def aplicar_cor_farol(
    celula,
    situacao,
):

    texto = str(
        situacao
    ).upper().strip()

    if "PRIORIT" in texto:

        celula.fill = FILL_PRIORITARIA

        celula.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color=COR_PRIORITARIA,
        )

    elif "ATEN" in texto:

        celula.fill = FILL_ATENCAO

        celula.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color=COR_ATENCAO,
        )

    elif "FAVOR" in texto:

        celula.fill = FILL_FAVORAVEL

        celula.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color=COR_FAVORAVEL,
        )

    else:

        celula.fill = FILL_SEM_HISTORICO

        celula.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color=COR_NEUTRO,
        )

    celula.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


# ==========================================================
# COR DOS GRUPOS DO CABEÇALHO
# ==========================================================

def aplicar_cores_cabecalho(
    ws,
):
    """
    Diferencia visualmente os grupos de avaliação
    utilizando exclusivamente tons de azul.
    """

    grupos = {

        "ADE": (
            [
                "LP_ABAIXO",
                "MAT_ABAIXO",
            ],
            COR_ADE,
        ),

        "PP1": (
            [
                "LP_PP1",
                "MAT_PP1",
            ],
            COR_PP1,
        ),

        "PP2": (
            [
                "LP_PP2",
                "MAT_PP2",
            ],
            COR_PP2,
        ),

        "PARTICIPACAO": (
            [
                "PART_ADE",
                "PART_PP1",
                "PART_PP2",
            ],
            COR_PARTICIPACAO,
        ),

        "EVOLUCAO": (
            [
                "EVOLUCAO_LP",
                "EVOLUCAO_MAT",
            ],
            COR_EVOLUCAO,
        ),

        "GERAL": (
            [
                "ESCOLA",
                "FAROL",
                "FAROL_URE",
                "COMPONENTE_FAROL",
                "DESTAQUE_EVOLUCAO",
                "DIAGNOSTICO",
                "ENCAMINHAMENTO",
            ],
            COR_GERAL,
        ),
    }

    for _, (
        colunas,
        cor,
    ) in grupos.items():

        preenchimento = PatternFill(
            fill_type="solid",
            fgColor=cor,
        )

        for nome_coluna in colunas:

            indice = obter_indice_coluna(
                ws,
                nome_coluna,
            )

            if indice is None:

                continue

            celula = ws.cell(
                row=1,
                column=indice,
            )

            celula.fill = preenchimento

            celula.font = FONTE_CABECALHO

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            celula.border = BORDA


# ==========================================================
# TEXTO DE DESTAQUE
# ==========================================================

def formatar_destaque(
    valor,
):

    if pd.isna(valor):

        return ""

    if isinstance(
        valor,
        bool,
    ):

        return (
            "★ Destaque"
            if valor
            else ""
        )

    texto = str(
        valor
    ).strip().upper()

    if texto in (
        "TRUE",
        "VERDADEIRO",
        "SIM",
        "1",
    ):

        return "★ Destaque"

    return ""


# ==========================================================
# FORMATA TEXTO LONGO
# ==========================================================

def formatar_texto_longo(
    ws,
    nome_coluna,
):

    indice = obter_indice_coluna(
        ws,
        nome_coluna,
    )

    if indice is None:

        return

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        celula = ws.cell(
            row=linha,
            column=indice,
        )

        celula.font = FONTE_TEXTO_LONGO

        celula.alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

        celula.border = BORDA


# ==========================================================
# ALTURA DA LINHA
# ==========================================================

def calcular_altura_linha(
    diagnostico,
    encaminhamento,
):

    texto_diagnostico = ""

    texto_encaminhamento = ""

    if pd.notna(
        diagnostico
    ):

        texto_diagnostico = str(
            diagnostico
        )

    if pd.notna(
        encaminhamento
    ):

        texto_encaminhamento = str(
            encaminhamento
        )

    maior_texto = max(
        len(texto_diagnostico),
        len(texto_encaminhamento),
    )

    texto_total = (
        len(texto_diagnostico)
        + len(texto_encaminhamento)
    )

    if maior_texto == 0:

        return 25

    if texto_total <= 120:

        return 30

    if texto_total <= 220:

        return 42

    if texto_total <= 320:

        return 54

    if texto_total <= 450:

        return 66

    return 66


# ==========================================================
# COR DA EVOLUÇÃO
# ==========================================================

def aplicar_cor_evolucao(
    celula,
):

    valor = celula.value

    if not isinstance(
        valor,
        (int, float),
    ):

        return

    if valor > 0:

        celula.font = Font(
            name="Aptos",
            size=9,
            bold=True,
            color=COR_EVOLUCAO_POSITIVA,
        )

    elif valor < 0:

        celula.font = Font(
            name="Aptos",
            size=9,
            bold=True,
            color=COR_EVOLUCAO_NEGATIVA,
        )

    else:

        celula.font = Font(
            name="Aptos",
            size=9,
            bold=True,
            color=COR_EVOLUCAO_NEUTRA,
        )


# ==========================================================
# RADAR PEDAGÓGICO
# ==========================================================

def criar_radar(
    writer,
    df,
):

    if df is None or df.empty:

        return

    # ======================================================
    # VALIDAÇÃO
    # ======================================================

    if "FAROL_URE" not in df.columns:

        raise ValueError(
            "A coluna FAROL_URE não foi encontrada."
        )

    # ======================================================
    # CÓPIA
    # ======================================================

    resultado = df.copy()

    # ======================================================
    # COLUNAS
    # ======================================================

    colunas_existentes = [
        coluna
        for coluna in COLUNAS_RADAR
        if coluna in resultado.columns
    ]

    resultado = resultado[
        colunas_existentes
    ].copy()

    # ======================================================
    # DESTAQUE
    # ======================================================

    if "DESTAQUE_EVOLUCAO" in resultado.columns:

        resultado[
            "DESTAQUE_EVOLUCAO"
        ] = (
            resultado[
                "DESTAQUE_EVOLUCAO"
            ].apply(
                formatar_destaque
            )
        )

    # ======================================================
    # FAROL
    # ======================================================

    indice_farol_ure = (
        resultado.columns.get_loc(
            "FAROL_URE"
        )
    )

    resultado.insert(
        indice_farol_ure,
        "FAROL",
        resultado[
            "FAROL_URE"
        ].apply(
            gerar_farol
        ),
    )

    # ======================================================
    # EXPORTAÇÃO
    # ======================================================

    resultado.to_excel(
        writer,
        sheet_name="RADAR PEDAGÓGICO",
        index=False,
    )

    ws = writer.sheets[
        "RADAR PEDAGÓGICO"
    ]

    # ======================================================
    # FORMATAÇÃO BASE
    # ======================================================

    formatar_planilha(
        ws
    )

    # ======================================================
    # CABEÇALHO
    # ======================================================

    aplicar_cores_cabecalho(
        ws
    )

    ws.row_dimensions[
        1
    ].height = 30

    # ======================================================
    # ÍNDICES
    # ======================================================

    indice_escola = obter_indice_coluna(
        ws,
        "ESCOLA",
    )

    indice_farol = obter_indice_coluna(
        ws,
        "FAROL",
    )

    indice_farol_ure = obter_indice_coluna(
        ws,
        "FAROL_URE",
    )

    indice_componente = obter_indice_coluna(
        ws,
        "COMPONENTE_FAROL",
    )

    indice_destaque = obter_indice_coluna(
        ws,
        "DESTAQUE_EVOLUCAO",
    )

    indice_diagnostico = obter_indice_coluna(
        ws,
        "DIAGNOSTICO",
    )

    indice_encaminhamento = obter_indice_coluna(
        ws,
        "ENCAMINHAMENTO",
    )

    indice_evolucao_lp = obter_indice_coluna(
        ws,
        "EVOLUCAO_LP",
    )

    indice_evolucao_mat = obter_indice_coluna(
        ws,
        "EVOLUCAO_MAT",
    )

    # ======================================================
    # CORPO
    # ======================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        for coluna in range(
            1,
            ws.max_column + 1,
        ):

            celula = ws.cell(
                row=linha,
                column=coluna,
            )

            celula.font = FONTE_NORMAL

            celula.border = BORDA

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=False,
            )

        # --------------------------------------------------
        # ESCOLA
        # --------------------------------------------------

        if indice_escola is not None:

            celula = ws.cell(
                row=linha,
                column=indice_escola,
            )

            celula.font = FONTE_ESCOLA

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # COMPONENTE
        # --------------------------------------------------

        if indice_componente is not None:

            celula = ws.cell(
                row=linha,
                column=indice_componente,
            )

            celula.font = FONTE_NORMAL

            celula.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )

        # --------------------------------------------------
        # DIAGNÓSTICO
        # --------------------------------------------------

        if indice_diagnostico is not None:

            celula = ws.cell(
                row=linha,
                column=indice_diagnostico,
            )

            celula.font = FONTE_TEXTO_LONGO

            celula.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )

            celula.fill = PatternFill(
                fill_type="solid",
                fgColor=COR_FUNDO_DIAGNOSTICO,
            )

        # --------------------------------------------------
        # ENCAMINHAMENTO
        # --------------------------------------------------

        if indice_encaminhamento is not None:

            celula = ws.cell(
                row=linha,
                column=indice_encaminhamento,
            )

            celula.font = FONTE_TEXTO_LONGO

            celula.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )

            celula.fill = PatternFill(
                fill_type="solid",
                fgColor=COR_FUNDO_ENCAMINHAMENTO,
            )

        # --------------------------------------------------
        # FAROL
        # --------------------------------------------------

        if (
            indice_farol is not None
            and indice_farol_ure is not None
        ):

            situacao = ws.cell(
                row=linha,
                column=indice_farol_ure,
            ).value

            aplicar_cor_farol(
                ws.cell(
                    row=linha,
                    column=indice_farol,
                ),
                situacao,
            )

            aplicar_cor_farol(
                ws.cell(
                    row=linha,
                    column=indice_farol_ure,
                ),
                situacao,
            )

        # --------------------------------------------------
        # DESTAQUE
        # --------------------------------------------------

        if indice_destaque is not None:

            celula = ws.cell(
                row=linha,
                column=indice_destaque,
            )

            if str(
                celula.value
            ).strip():

                celula.fill = FILL_DESTAQUE

                celula.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="1F4E79",
                )

                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        # --------------------------------------------------
        # EVOLUÇÃO LP
        # --------------------------------------------------

        if indice_evolucao_lp is not None:

            aplicar_cor_evolucao(
                ws.cell(
                    row=linha,
                    column=indice_evolucao_lp,
                )
            )

        # --------------------------------------------------
        # EVOLUÇÃO MAT
        # --------------------------------------------------

        if indice_evolucao_mat is not None:

            aplicar_cor_evolucao(
                ws.cell(
                    row=linha,
                    column=indice_evolucao_mat,
                )
            )

    # ======================================================
    # PERCENTUAIS
    # ======================================================

    colunas_percentuais = [
        "LP_ABAIXO",
        "MAT_ABAIXO",
        "LP_PP1",
        "MAT_PP1",
        "LP_PP2",
        "MAT_PP2",
        "PART_ADE",
        "PART_PP1",
        "PART_PP2",
        "EVOLUCAO_LP",
        "EVOLUCAO_MAT",
    ]

    for nome_coluna in colunas_percentuais:

        indice = obter_indice_coluna(
            ws,
            nome_coluna,
        )

        if indice is None:

            continue

        for linha in range(
            2,
            ws.max_row + 1,
        ):

            celula = ws.cell(
                row=linha,
                column=indice,
            )

            celula.number_format = "0.0%"

            if nome_coluna not in (
                "EVOLUCAO_LP",
                "EVOLUCAO_MAT",
            ):

                celula.font = Font(
                    name="Aptos",
                    size=9,
                    bold=True,
                    color="000000",
                )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # ======================================================
    # REAPLICA COR DA EVOLUÇÃO
    # ======================================================

    if indice_evolucao_lp is not None:

        for linha in range(
            2,
            ws.max_row + 1,
        ):

            aplicar_cor_evolucao(
                ws.cell(
                    row=linha,
                    column=indice_evolucao_lp,
                )
            )

    if indice_evolucao_mat is not None:

        for linha in range(
            2,
            ws.max_row + 1,
        ):

            aplicar_cor_evolucao(
                ws.cell(
                    row=linha,
                    column=indice_evolucao_mat,
                )
            )

    # ======================================================
    # LARGURAS
    # ======================================================

    definir_largura_coluna(
        ws,
        "ESCOLA",
        32,
    )

    definir_largura_coluna(
        ws,
        "FAROL",
        6,
    )

    definir_largura_coluna(
        ws,
        "FAROL_URE",
        15,
    )

    definir_largura_coluna(
        ws,
        "COMPONENTE_FAROL",
        20,
    )

    definir_largura_coluna(
        ws,
        "DESTAQUE_EVOLUCAO",
        16,
    )

    definir_largura_coluna(
        ws,
        "DIAGNOSTICO",
        46,
    )

    definir_largura_coluna(
        ws,
        "ENCAMINHAMENTO",
        50,
    )

    # ======================================================
    # TEXTOS LONGOS
    # ======================================================

    formatar_texto_longo(
        ws,
        "DIAGNOSTICO",
    )

    formatar_texto_longo(
        ws,
        "ENCAMINHAMENTO",
    )

    # ======================================================
    # ALTURA DAS LINHAS
    # ======================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        diagnostico = ""

        encaminhamento = ""

        if (
            "DIAGNOSTICO"
            in resultado.columns
        ):

            diagnostico = resultado.iloc[
                linha - 2
            ]["DIAGNOSTICO"]

        if (
            "ENCAMINHAMENTO"
            in resultado.columns
        ):

            encaminhamento = resultado.iloc[
                linha - 2
            ]["ENCAMINHAMENTO"]

        altura = calcular_altura_linha(
            diagnostico,
            encaminhamento,
        )

        ws.row_dimensions[
            linha
        ].height = altura

    # ======================================================
    # SEPARAÇÃO VISUAL ENTRE ESCOLAS
    # ======================================================

    for linha in range(
        2,
        ws.max_row + 1,
    ):

        for coluna in range(
            1,
            ws.max_column + 1,
        ):

            celula = ws.cell(
                row=linha,
                column=coluna,
            )

            if linha == ws.max_row:

                celula.border = (
                    BORDA_SEPARACAO_ESCOLA
                )

                continue

            escola_atual = ws.cell(
                row=linha,
                column=indice_escola,
            ).value

            escola_seguinte = ws.cell(
                row=linha + 1,
                column=indice_escola,
            ).value

            if escola_atual != escola_seguinte:

                celula.border = (
                    BORDA_SEPARACAO_ESCOLA
                )

    # ======================================================
    # FINALIZAÇÃO
    # ======================================================

    ws.freeze_panes = "A2"

    if ws.max_row >= 2:

        ws.auto_filter.ref = ws.dimensions

    ws.sheet_view.zoomScale = 90

    ws.sheet_view.showGridLines = False
