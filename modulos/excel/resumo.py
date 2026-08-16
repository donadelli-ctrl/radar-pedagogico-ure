"""
==========================================================
RADAR PEDAGÓGICO URE
MÓDULO: resumo.py
Versão: 5.0
==========================================================

Responsabilidade:
Gerar a aba RESUMO EXECUTIVO.

A aba apresenta uma visão gerencial do Radar URE,
com identidade visual clara, moderna e de alto contraste.

A formatação não altera os critérios pedagógicos
nem os cálculos dos indicadores.
"""


from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)


# ==========================================================
# CORES DA IDENTIDADE VISUAL
# ==========================================================

# Fundo do título
COR_TITULO = "D9EDF5"

# Texto do título — preto
COR_TITULO_FONTE = "111111"

# Subtítulo — preto/cinza escuro
COR_SUBTITULO = "333333"


# ==========================================================
# CORES DOS CARTÕES
# ==========================================================

COR_TOTAL = "C7E3F0"

COR_PRIORITARIA = "F2A3A3"

COR_ATENCAO = "FFD966"

COR_FAVORAVEL = "A9D18E"

COR_DESTAQUE = "8FBCE6"

COR_SEM_DADOS = "D9D9D9"


# ==========================================================
# CORES DOS TEXTOS
# ==========================================================

COR_TEXTO = "222222"

COR_TEXTO_FORTE = "111111"


# ==========================================================
# CORES DAS BORDAS E FUNDOS
# ==========================================================

COR_BORDA = "9FBAC5"

COR_FUNDO = "FFFFFF"

COR_FUNDO_BLOCO = "F4F7F9"


# ==========================================================
# CORES DOS ÍCONES
# ==========================================================

COR_ICONE_TOTAL = "1976A8"

COR_ICONE_PRIORITARIA = "C62828"

COR_ICONE_ATENCAO = "D68B00"

COR_ICONE_FAVORAVEL = "3F7D3F"

COR_ICONE_DESTAQUE = "246B91"


# ==========================================================
# FONTES
# ==========================================================

FONTE_TITULO = Font(
    name="Aptos",
    size=20,
    bold=True,
    color=COR_TITULO_FONTE,
)


FONTE_SUBTITULO = Font(
    name="Aptos",
    size=11,
    bold=True,
    color=COR_TITULO_FONTE,
)


FONTE_TEXTO = Font(
    name="Aptos",
    size=10,
    color=COR_TEXTO,
)


FONTE_NUMERO = Font(
    name="Aptos",
    size=20,
    bold=True,
    color=COR_TEXTO_FORTE,
)


FONTE_RODAPE = Font(
    name="Aptos",
    size=9,
    italic=True,
    color="666666",
)


# ==========================================================
# ALINHAMENTOS
# ==========================================================

CENTRO = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)


ESQUERDA = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True,
)


DIREITA = Alignment(
    horizontal="right",
    vertical="center",
)


# ==========================================================
# BORDA
# ==========================================================

BORDA = Border(
    left=Side(
        style="thin",
        color=COR_BORDA,
    ),
    right=Side(
        style="thin",
        color=COR_BORDA,
    ),
    top=Side(
        style="thin",
        color=COR_BORDA,
    ),
    bottom=Side(
        style="thin",
        color=COR_BORDA,
    ),
)


# ==========================================================
# FUNÇÃO PARA CRIAR CARTÃO
# ==========================================================

def criar_cartao(
    ws,
    col_inicio,
    linha,
    icone,
    titulo,
    valor,
    cor_fundo,
    cor_icone,
):
    """
    Cria um cartão visual de indicador.

    O cartão possui:
    - ícone;
    - título;
    - número.

    A função trabalha somente com apresentação.
    """

    coluna_numero = (
        ord(col_inicio)
        - ord("A")
        + 1
    )

    col_fim_numero = (
        coluna_numero
        + 1
    )

    col_fim = chr(
        ord("A")
        + col_fim_numero
        - 1
    )

    # ======================================================
    # MESCLAGEM
    # ======================================================

    ws.merge_cells(
        f"{col_inicio}{linha}:"
        f"{col_fim}{linha}"
    )

    ws.merge_cells(
        f"{col_inicio}{linha + 1}:"
        f"{col_fim}{linha + 1}"
    )

    ws.merge_cells(
        f"{col_inicio}{linha + 2}:"
        f"{col_fim}{linha + 2}"
    )

    # ======================================================
    # CÉLULAS
    # ======================================================

    c1 = ws[
        f"{col_inicio}{linha}"
    ]

    c2 = ws[
        f"{col_inicio}{linha + 1}"
    ]

    c3 = ws[
        f"{col_inicio}{linha + 2}"
    ]

    c1.value = icone
    c2.value = titulo
    c3.value = valor

    # ======================================================
    # ÍCONE
    # ======================================================

    c1.font = Font(
        name="Arial",
        size=16,
        bold=True,
        color=cor_icone,
    )

    c1.alignment = CENTRO

    # ======================================================
    # TÍTULO
    # ======================================================

    c2.font = Font(
        name="Aptos",
        size=10,
        bold=True,
        color=COR_TEXTO_FORTE,
    )

    c2.alignment = CENTRO

    # ======================================================
    # NÚMERO
    # ======================================================

    c3.font = FONTE_NUMERO

    c3.alignment = CENTRO

    # ======================================================
    # PREENCHIMENTO
    # ======================================================

    preenchimento = PatternFill(
        fill_type="solid",
        fgColor=cor_fundo,
    )

    for r in range(
        linha,
        linha + 3,
    ):

        for coluna in (
            col_inicio,
            col_fim,
        ):

            celula = ws[
                f"{coluna}{r}"
            ]

            celula.fill = preenchimento
            celula.border = BORDA


# ==========================================================
# RESUMO EXECUTIVO
# ==========================================================

def criar_resumo(
    writer,
    indicadores,
):
    """
    Cria a aba RESUMO EXECUTIVO.

    Os dados são provenientes do dicionário
    de indicadores já calculado.
    """

    ws = writer.book.create_sheet(
        "RESUMO EXECUTIVO"
    )

    # ======================================================
    # CONFIGURAÇÃO
    # ======================================================

    ws.sheet_view.showGridLines = False

    ws.freeze_panes = "A5"

    # ======================================================
    # LARGURA DAS COLUNAS
    # ======================================================

    larguras = {
        "A": 18,
        "B": 18,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 18,
        "H": 18,
    }

    for coluna, largura in larguras.items():

        ws.column_dimensions[
            coluna
        ].width = largura

    # ======================================================
    # ALTURA DAS LINHAS
    # ======================================================

    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 25

    # ======================================================
    # TÍTULO
    # ======================================================

    ws.merge_cells(
        "A1:H1"
    )

    titulo = ws["A1"]

    titulo.value = (
        "▮ RADAR PEDAGÓGICO URE"
    )

    titulo.font = Font(
        name="Aptos",
        size=20,
        bold=True,
        color="111111",
    )

    titulo.fill = PatternFill(
        fill_type="solid",
        fgColor=COR_TITULO,
    )

    titulo.alignment = CENTRO

    # ======================================================
    # SUBTÍTULO
    # ======================================================

    ws.merge_cells(
        "A2:H2"
    )

    subtitulo = ws["A2"]

    subtitulo.value = (
        "Painel Executivo das Escolas da URE"
    )

    subtitulo.font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color="333333",
    )

    subtitulo.alignment = CENTRO

    # ======================================================
    # INFORMAÇÕES DA GERAÇÃO
    # ======================================================

    ws["A4"] = "Data de geração"

    ws["A4"].font = Font(
        name="Aptos",
        size=10,
        bold=True,
        color="111111",
    )

    ws["B4"] = indicadores[
        "DATA_GERACAO"
    ].strftime(
        "%d/%m/%Y %H:%M"
    )

    ws["B4"].font = FONTE_TEXTO

    ws["E4"] = "Avaliações"

    ws["E4"].font = Font(
        name="Aptos",
        size=10,
        bold=True,
        color="111111",
    )

    ws["F4"] = ", ".join(
        indicadores[
            "AVALIACOES"
        ]
    )

    ws["F4"].font = FONTE_TEXTO

    # ======================================================
    # CARTÕES PRINCIPAIS
    # ======================================================

    criar_cartao(
        ws,
        "A",
        6,
        "●",
        "TOTAL DE ESCOLAS",
        indicadores[
            "TOTAL_ESCOLAS"
        ],
        COR_TOTAL,
        COR_ICONE_TOTAL,
    )

    criar_cartao(
        ws,
        "C",
        6,
        "●",
        "PRIORITÁRIAS",
        indicadores[
            "PRIORITARIA"
        ],
        COR_PRIORITARIA,
        COR_ICONE_PRIORITARIA,
    )

    criar_cartao(
        ws,
        "E",
        6,
        "●",
        "ATENÇÃO",
        indicadores[
            "ATENCAO"
        ],
        COR_ATENCAO,
        COR_ICONE_ATENCAO,
    )

    criar_cartao(
        ws,
        "G",
        6,
        "●",
        "FAVORÁVEIS",
        indicadores.get(
            "FAVORAVEL",
            indicadores.get(
                "FAVORAVEIS",
                0,
            ),
        ),
        COR_FAVORAVEL,
        COR_ICONE_FAVORAVEL,
    )

    # ======================================================
    # DESTAQUE DE EVOLUÇÃO
    # ======================================================

    criar_cartao(
        ws,
        "A",
        11,
        "★",
        "DESTAQUE DE EVOLUÇÃO",
        indicadores.get(
            "DESTAQUE",
            0,
        ),
        COR_DESTAQUE,
        COR_ICONE_DESTAQUE,
    )

    # ======================================================
    # PARTICIPAÇÃO
    # ======================================================

    ws["E11"] = (
        "◉ PARTICIPAÇÃO"
    )

    ws["E11"].font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color=COR_TEXTO_FORTE,
    )

    ws["E11"].alignment = ESQUERDA

    participacao = [
        (
            "ADE",
            indicadores.get(
                "PART_ADE"
            ),
        ),
        (
            "PP1",
            indicadores.get(
                "PART_PP1"
            ),
        ),
        (
            "PP2",
            indicadores.get(
                "PART_PP2"
            ),
        ),
    ]

    linha = 12

    for nome, valor in participacao:

        ws[f"E{linha}"] = nome

        ws[f"F{linha}"] = valor

        ws[
            f"E{linha}"
        ].font = FONTE_TEXTO

        ws[
            f"F{linha}"
        ].font = FONTE_NUMERO

        ws[
            f"F{linha}"
        ].alignment = DIREITA

        if valor is not None:

            ws[
                f"F{linha}"
            ].number_format = "0.0%"

        ws[
            f"E{linha}"
        ].border = BORDA

        ws[
            f"F{linha}"
        ].border = BORDA

        preenchimento = PatternFill(
            fill_type="solid",
            fgColor=COR_FUNDO_BLOCO,
        )

        ws[
            f"E{linha}"
        ].fill = preenchimento

        ws[
            f"F{linha}"
        ].fill = preenchimento

        linha += 1

    # ======================================================
    # MÉDIAS GERAIS
    # ======================================================

    ws["A16"] = (
        "▣ MÉDIAS GERAIS"
    )

    ws["A16"].font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color=COR_TEXTO_FORTE,
    )

    ws["A16"].alignment = ESQUERDA

    medias = [
        (
            "PP1",
            indicadores.get(
                "MEDIA_PP1"
            ),
        ),
        (
            "PP2",
            indicadores.get(
                "MEDIA_PP2"
            ),
        ),
    ]

    linha = 17

    for nome, valor in medias:

        ws[f"A{linha}"] = nome

        ws[f"B{linha}"] = valor

        ws[
            f"A{linha}"
        ].font = FONTE_TEXTO

        ws[
            f"B{linha}"
        ].font = FONTE_NUMERO

        ws[
            f"B{linha}"
        ].alignment = DIREITA

        if valor is not None:

            ws[
                f"B{linha}"
            ].number_format = "0.0%"

        ws[
            f"A{linha}"
        ].border = BORDA

        ws[
            f"B{linha}"
        ].border = BORDA

        preenchimento = PatternFill(
            fill_type="solid",
            fgColor=COR_FUNDO_BLOCO,
        )

        ws[
            f"A{linha}"
        ].fill = preenchimento

        ws[
            f"B{linha}"
        ].fill = preenchimento

        linha += 1

    # ======================================================
    # CLASSIFICAÇÃO DO FAROL
    # ======================================================

    ws["E16"] = (
        "● CLASSIFICAÇÃO DO FAROL"
    )

    ws["E16"].font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color=COR_TEXTO_FORTE,
    )

    ws["E16"].alignment = ESQUERDA

    classificacao = [
        (
            "●  Prioritárias",
            indicadores[
                "PRIORITARIA"
            ],
            COR_PRIORITARIA,
            COR_ICONE_PRIORITARIA,
        ),
        (
            "●  Atenção",
            indicadores[
                "ATENCAO"
            ],
            COR_ATENCAO,
            COR_ICONE_ATENCAO,
        ),
        (
            "●  Favoráveis",
            indicadores.get(
                "FAVORAVEL",
                indicadores.get(
                    "FAVORAVEIS",
                    0,
                ),
            ),
            COR_FAVORAVEL,
            COR_ICONE_FAVORAVEL,
        ),
        (
            "●  Sem dados",
            indicadores[
                "SEM_DADOS"
            ],
            COR_SEM_DADOS,
            "555555",
        ),
    ]

    linha = 17

    for (
        texto,
        valor,
        cor,
        cor_texto,
    ) in classificacao:

        ws[
            f"E{linha}"
        ] = texto

        ws[
            f"F{linha}"
        ] = valor

        ws[
            f"E{linha}"
        ].font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color="111111",
        )

        ws[
            f"F{linha}"
        ].font = Font(
            name="Aptos",
            size=16,
            bold=True,
            color="111111",
        )

        ws[
            f"F{linha}"
        ].alignment = DIREITA

        preenchimento = PatternFill(
            fill_type="solid",
            fgColor=cor,
        )

        ws[
            f"E{linha}"
        ].fill = preenchimento

        ws[
            f"F{linha}"
        ].fill = preenchimento

        ws[
            f"E{linha}"
        ].border = BORDA

        ws[
            f"F{linha}"
        ].border = BORDA

        linha += 1

    # ======================================================
    # RODAPÉ
    # ======================================================

    ws.merge_cells(
        "A24:H24"
    )

    rodape = ws["A24"]

    rodape.value = (
        "RADAR PEDAGÓGICO URE • "
        "Painel Executivo"
    )

    rodape.font = FONTE_RODAPE

    rodape.alignment = CENTRO

    # ======================================================
    # LINHA DE SEPARAÇÃO
    # ======================================================

    for coluna in range(
        1,
        9,
    ):

        ws.cell(
            row=10,
            column=coluna,
        ).border = Border(
            bottom=Side(
                style="medium",
                color=COR_BORDA,
            )
        )

    # ======================================================
    # VISUALIZAÇÃO
    # ======================================================

    ws.sheet_view.zoomScale = 90

    ws.sheet_view.showGridLines = False